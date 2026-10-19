"""Scenario tests for the chat panel on the Models tab: saved conversations, the search index, streaming replies from
Ollama's /api/chat, and file attachments.

Run alone:  .venv/bin/python -m unittest discover -s tests -p test_scenarios_chat.py
Nothing here touches real data, Ollama or the network: data folders point at a temporary directory (TempDataTest)
and every HTTP call is replaced with a fake."""

import base64
import json
import re
import sys
import time
import types
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_app import FakeResponse, TempDataTest, app  # noqa: E402


def ev(content="", done=False, **extra):
    """One ndjson line of a streamed /api/chat reply."""
    d = {"message": {"role": "assistant", "content": content}}
    if done:
        d["done"] = True
    d.update(extra)
    return json.dumps(d).encode()


def make_chat(cid="c1", model="m", messages=None, title="t"):
    return {"id": cid, "title": title, "model": model, "created": 1.0, "totals": app.new_totals(),
            "messages": messages if messages is not None else [{"role": "user", "content": "hi"}]}


def is_noop(update):
    return update == {"__type__": "update"}


# ---------------------------------------------------------------- plain helpers
class ContentAndTotals(unittest.TestCase):
    def test_content_text_joins_a_tuple_of_strings(self):
        self.assertEqual(app.content_text(("ab", "cd", "")), "abcd")

    def test_content_text_flattens_nested_lists(self):
        self.assertEqual(app.content_text(["a", ["b", [{"text": "c"}]], {"content": "d"}]), "abcd")

    def test_content_text_dict_prefers_text_then_falls_back_to_content(self):
        self.assertEqual(app.content_text({"text": "T", "content": "C"}), "T")
        self.assertEqual(app.content_text({"text": "", "content": "C"}), "C")
        self.assertEqual(app.content_text({"path": "/x.png"}), "")
        self.assertEqual(app.content_text(0), "0")  # a zero is content, not nothing

    def test_new_totals_returns_an_independent_dict_each_time(self):
        a, b = app.new_totals(), app.new_totals()
        a["replies"] = 7
        self.assertEqual(b, {"replies": 0, "seconds": 0.0, "prompt": 0, "output": 0})

    def test_empty_chat_is_unsaved_and_carries_the_model(self):
        c1, c2 = app.empty_chat("llama3"), app.empty_chat()
        self.assertEqual((c1["id"], c1["model"], c1["messages"], c1["created"], c1["title"]), (None, "llama3", [], None, ""))
        self.assertIsNone(c2["model"])
        c1["messages"].append("x")
        c1["totals"]["replies"] = 3
        self.assertEqual((c2["messages"], c2["totals"]["replies"]), ([], 0))

    def test_totals_line_uses_singular_reply(self):
        md = app.chat_totals_md({"replies": 1, "seconds": 2.5, "prompt": 3, "output": 4})
        self.assertIn("1 reply ", md)
        self.assertIn("2.5 s", md)
        self.assertIn("7 total", md)

    def test_totals_line_plural_with_thousands_separators_and_minutes(self):
        md = app.chat_totals_md({"replies": 12, "seconds": 125.0, "prompt": 12345, "output": 987654})
        self.assertIn("12 replies", md)
        self.assertIn("2 min 5 s", md)
        self.assertIn("12,345 prompt tokens", md)
        self.assertIn("987,654 response tokens", md)
        self.assertIn("999,999 total", md)

    def test_totals_line_without_replies_says_no_messages_even_with_tokens(self):
        self.assertEqual(app.chat_totals_md({"replies": 0, "seconds": 9.0, "prompt": 50, "output": 5}), "No messages yet.")


FMT_DURATION_CASES = [
    ("zero", 0, "0.0 s"), ("ninety_seconds", 90, "1 min 30 s"),
    ("half_second_rounds_to_two_minutes", 119.5, "2 min 0 s"), ("just_under_an_hour_rounds_up", 3599.6, "1 h 0 min"),
    ("exactly_one_hour", 3600, "1 h 0 min"), ("two_hours_two_minutes", 7322, "2 h 2 min"),
    ("a_whole_day", 86400, "24 h 0 min"),
]


class FmtDuration(unittest.TestCase):
    pass


for _name, _sec, _want in FMT_DURATION_CASES:
    def _t(self, sec=_sec, want=_want):
        self.assertEqual(app._fmt_duration(sec), want)
    setattr(FmtDuration, f"test_fmt_duration_{_name}", _t)


CHAT_PATH_CASES = [
    ("plain_word", "abc", True), ("generated_id", "20261019-131047-a1b2c3", True), ("underscore", "with_under_score", True),
    ("unicode_letters", "日本語", True), ("space", "a b", False), ("semicolon", "a;b", False), ("dot_dot", "..", False),
    ("nul_byte", "x\x00", False),
]


class ChatPath(TempDataTest):
    pass


for _name, _cid, _ok in CHAT_PATH_CASES:
    def _t(self, cid=_cid, ok=_ok):
        path = app._chat_path(cid)
        if ok:
            self.assertEqual(path, app.CHATS_DIR / f"{cid}.json")
        else:
            self.assertIsNone(path)
    setattr(ChatPath, f"test_chat_path_{_name}_{'accepted' if _ok else 'rejected'}", _t)


# ---------------------------------------------------------------- search index
def idx_chat(cid, title, *messages, updated=0, role="user"):
    return {"id": cid, "title": title, "updated": updated, "messages": [{"role": role, "content": m} for m in messages]}


def search_corpus():
    idx = app.ChatIndex()
    idx.add(idx_chat("c1", "Python decorators", "How do decorators work in Python 3.12?", updated=1))
    c2 = idx_chat("c2", "Cooking pasta", "Carbonara recipe with eggs", updated=2)
    c2["messages"].append({"role": "assistant", "content": "Use guanciale and pecorino"})
    idx.add(c2)
    idx.add(idx_chat("c3", "Trip to Paris", "Best museums: Louvre, Orsay", updated=3))
    idx.add(idx_chat("c4", "E-mail draft", "Write an e-mail to my boss about the Q3 report", updated=4))
    idx.add(idx_chat("c5", "", "python snake facts", updated=5))
    c6 = idx_chat("c6", "Security scan", "garak probes for prompt injection", updated=6)
    c6["messages"][0]["attachments"] = [{"name": "report.csv", "kind": "text", "text": "dan jailbreak results"}]
    idx.add(c6)
    return idx


SEARCH_CASES = [
    ("lowercase_word_title_and_body", "python", ["c5", "c1"]),
    ("uppercase_query_is_case_insensitive", "PYTHON", ["c5", "c1"]),
    ("middle_of_word_does_not_match", "thon", []),
    ("terms_in_any_order", "decorators python", ["c1"]),
    ("all_terms_must_match", "python snake", ["c5"]),
    ("assistant_reply_is_searchable", "guanciale", ["c2"]),
    ("title_only_word", "pasta", ["c2"]),
    ("terms_split_by_punctuation_in_body", "louvre orsay", ["c3"]),
    ("terms_from_different_chats_match_nothing", "louvre pasta", []),
    ("hyphenated_query", "e-mail", ["c4"]),
    ("letter_digit_word", "q3", ["c4"]),
    ("attachment_text", "jailbreak", ["c6"]),
    ("attachment_name_and_body_word_newest_first", "report", ["c6", "c4"]),
    ("messy_whitespace_and_punctuation", "  garak   ,probes!! ", ["c6"]),
    ("empty_query_lists_everything_newest_first", "", ["c6", "c5", "c4", "c3", "c2", "c1"]),
]


class ChatSearchQueries(unittest.TestCase):
    pass


for _name, _query, _want in SEARCH_CASES:
    def _t(self, query=_query, want=_want):
        self.assertEqual([c["id"] for c in search_corpus().search(query)], want)
    setattr(ChatSearchQueries, f"test_search_{_name}", _t)


class ChatIndexScenarios(TempDataTest):
    # BUG: ChatIndex.WORD only knows [a-z0-9], so a query in Cyrillic/CJK has no terms and search() returns every
    # chat instead of only the ones containing the word (app.py ChatIndex.WORD / search, ~lines 955 and 1010).
    def test_non_latin_query_only_finds_chats_containing_it(self):
        idx = app.ChatIndex()
        idx.add(idx_chat("ru", "Привет", "привет мир", updated=1))
        idx.add(idx_chat("en", "Hello", "hello world", updated=2))
        self.assertEqual([c["id"] for c in idx.search("привет")], ["ru"])

    def test_accented_query_finds_the_accented_chat(self):
        idx = app.ChatIndex()
        idx.add(idx_chat("a", "Lunch", "meet at the café", updated=1))
        idx.add(idx_chat("b", "Other", "nothing related", updated=2))
        self.assertEqual([c["id"] for c in idx.search("café")], ["a"])

    def test_punctuation_only_query_lists_all_chats(self):
        idx = search_corpus()
        self.assertEqual(len(idx.search("?!... ---")), 6)
        self.assertEqual(len(idx.search(None)), 6)

    def test_blank_title_is_listed_as_new_chat(self):
        idx = search_corpus()
        self.assertEqual(idx.search("snake")[0]["title"], "New chat")

    def test_touch_moves_a_chat_to_the_top_and_ignores_unknown_chats(self):
        idx = search_corpus()
        idx.touch({"id": "c1", "updated": 100})
        idx.touch({"id": "ghost", "updated": 200})
        ids = [c["id"] for c in idx.search("")]
        self.assertEqual(ids[0], "c1")
        self.assertNotIn("ghost", ids)

    def test_build_reads_saved_chats_and_skips_temp_and_corrupt_files(self):
        app.CHATS_DIR.mkdir()
        for cid, word, upd in (("x1", "alpha", 1), ("x2", "bravo", 2)):
            (app.CHATS_DIR / f"{cid}.json").write_text(json.dumps(idx_chat(cid, word, "shared", updated=upd)))
        (app.CHATS_DIR / "x3.tmp").write_text(json.dumps(idx_chat("x3", "charlie", "shared")))
        (app.CHATS_DIR / "x4.json").write_text("")
        app.CHAT_INDEX.build()
        self.assertEqual([c["id"] for c in app.CHAT_INDEX.search("shared")], ["x2", "x1"])
        self.assertEqual(app.CHAT_INDEX.search("charlie"), [])

    def test_build_without_a_chats_folder_is_harmless(self):
        self.assertFalse(app.CHATS_DIR.exists())
        app.CHAT_INDEX.build()
        self.assertEqual(app.CHAT_INDEX.search(""), [])

    def test_removing_one_chat_keeps_a_shared_word_for_the_other(self):
        idx = app.ChatIndex()
        idx.add(idx_chat("a", "x", "kubernetes cluster", updated=1))
        idx.add(idx_chat("b", "y", "kubernetes pods", updated=2))
        idx.search("kube")
        idx.remove("b")
        self.assertEqual([c["id"] for c in idx.search("kube")], ["a"])
        self.assertEqual(idx.search("pods"), [])

    def test_remove_unknown_chat_leaves_index_intact(self):
        idx = search_corpus()
        idx.remove("does-not-exist")
        idx.remove(None)
        self.assertEqual(len(idx.search("")), 6)

    def test_gradio_list_content_is_indexed(self):
        idx = app.ChatIndex()
        idx.add({"id": "g", "title": "", "updated": 1,
                 "messages": [{"role": "user", "content": [{"type": "text", "text": "photosynthesis basics"}]}]})
        self.assertEqual([c["id"] for c in idx.search("photosyn")], ["g"])


# ---------------------------------------------------------------- saving, loading, deleting
class ChatStorage(TempDataTest):
    def test_save_sets_updated_and_leaves_no_temp_file(self):
        before = time.time()
        chat = make_chat("s1")
        app.save_chat(chat)
        self.assertGreaterEqual(chat["updated"], before)
        self.assertEqual(sorted(p.name for p in app.CHATS_DIR.iterdir()), ["s1.json"])

    def test_save_without_index_writes_the_file_but_does_not_index_a_new_chat(self):
        app.save_chat(make_chat("s2", title="Streaming"), index=False)
        self.assertEqual(app.load_chat("s2")["title"], "Streaming")
        self.assertEqual(app.CHAT_INDEX.search("streaming"), [])

    def test_save_without_index_refreshes_order_but_keeps_old_words(self):
        a, b = make_chat("a", title="apple"), make_chat("b", title="banana")
        app.save_chat(a)
        app.save_chat(b)
        a["title"] = "cherry"
        a["updated"] = 0
        app.save_chat(a, index=False)
        self.assertEqual([c["id"] for c in app.CHAT_INDEX.search("")], ["a", "b"])
        self.assertEqual(app.CHAT_INDEX.search("cherry"), [])
        self.assertEqual([c["id"] for c in app.CHAT_INDEX.search("apple")], ["a"])

    def test_unicode_title_and_messages_round_trip(self):
        chat = make_chat("u1", title="Вопрос про 日本 🍣", messages=[{"role": "user", "content": "naïve façade — 😀"}])
        app.save_chat(chat)
        loaded = app.load_chat("u1")
        self.assertEqual(loaded["title"], "Вопрос про 日本 🍣")
        self.assertEqual(loaded["messages"][0]["content"], "naïve façade — 😀")

    def test_very_long_title_round_trips_and_is_searchable(self):
        title = "word " * 2000 + "needle"
        app.save_chat(make_chat("long", title=title))
        self.assertEqual(app.load_chat("long")["title"], title)
        self.assertEqual([c["id"] for c in app.CHAT_INDEX.search("needle")], ["long"])

    def test_load_missing_chat_returns_none(self):
        app.CHATS_DIR.mkdir()
        self.assertIsNone(app.load_chat("never-saved"))

    def test_load_refuses_paths_outside_the_chats_folder(self):
        (app.DATA_DIR / "secret.json").write_text(json.dumps({"id": "secret"}))
        self.assertIsNone(app.load_chat("../secret"))

    def test_delete_of_a_chat_without_file_still_drops_it_from_search(self):
        app.CHAT_INDEX.add(idx_chat("orphan", "orphaned chat"))
        self.assertTrue(app.delete_chat_file("orphan"))
        self.assertEqual(app.CHAT_INDEX.search("orphaned"), [])

    def test_delete_only_removes_that_chats_attachments(self):
        for cid in ("d1", "d2"):
            app.save_chat(make_chat(cid))
            (app.ATTACH_DIR / cid).mkdir(parents=True)
            (app.ATTACH_DIR / cid / "f.txt").write_text("x")
        app.delete_chat_file("d1")
        self.assertFalse((app.ATTACH_DIR / "d1").exists())
        self.assertTrue((app.ATTACH_DIR / "d2" / "f.txt").exists())
        self.assertIsNotNone(app.load_chat("d2"))

    def test_resaving_replaces_contents_and_search_words(self):
        chat = make_chat("r1", messages=[{"role": "user", "content": "first version"}])
        app.save_chat(chat)
        chat["messages"] = [{"role": "user", "content": "second edition"}]
        app.save_chat(chat)
        self.assertEqual(app.load_chat("r1")["messages"][0]["content"], "second edition")
        self.assertEqual(app.CHAT_INDEX.search("version"), [])
        self.assertEqual(len(app.CHAT_INDEX.search("edition")), 1)


# ---------------------------------------------------------------- display and model lock
class DisplayAndLock(unittest.TestCase):
    def test_display_messages_of_nothing_is_empty(self):
        self.assertEqual(app.display_messages(None), [])
        self.assertEqual(app.display_messages({}), [])

    def test_display_adds_stats_under_assistant_replies_only(self):
        chat = {"messages": [{"role": "user", "content": "q", "stats": "ignored"},
                             {"role": "assistant", "content": "a", "stats": "1.0 s"},
                             {"role": "assistant", "content": "b", "stats": ""}]}
        out = app.display_messages(chat)
        self.assertEqual([m["content"] for m in out], ["q", "a\n\n<sub>1.0 s</sub>", "b"])
        self.assertEqual([m["role"] for m in out], ["user", "assistant", "assistant"])

    def test_display_puts_attachments_before_the_text(self):
        chat = {"messages": [{"role": "user", "content": "see this",
                              "attachments": [{"kind": "text", "name": "a.txt", "size": 10, "path": "/x/a.txt"}]}]}
        text = app.display_messages(chat)[0]["content"]
        self.assertTrue(text.startswith('<div class="attach-row">'))
        self.assertTrue(text.endswith("see this"))

    def test_model_lock_for_no_chat_is_free(self):
        with mock.patch.object(app, "model_names", side_effect=AssertionError("not needed")):
            up = app.model_lock(None)
        self.assertEqual((up["interactive"], up["label"]), (True, "Model"))

    def test_model_lock_locks_to_an_installed_model(self):
        with mock.patch.object(app, "model_names", return_value=["a", "m"]):
            up = app.model_lock(make_chat())
        self.assertEqual((up["value"], up["choices"], up["interactive"]), ("m", ["a", "m"], False))
        self.assertIn("locked", up["label"])

    def test_model_lock_keeps_a_since_deleted_model_selectable(self):
        with mock.patch.object(app, "model_names", return_value=["a"]):
            up = app.model_lock(make_chat(model="gone:7b"))
        self.assertEqual(up["choices"], ["a", "gone:7b"])
        self.assertEqual(up["value"], "gone:7b")

    def test_model_lock_without_model_or_messages_stays_free(self):
        with mock.patch.object(app, "model_names", return_value=["a"]):
            self.assertTrue(app.model_lock(make_chat(model=None))["interactive"])
            self.assertTrue(app.model_lock(make_chat(messages=[]))["interactive"])


# ---------------------------------------------------------------- sending a message
class ChatSubmit(TempDataTest):
    def setUp(self):
        super().setUp()
        p = mock.patch.object(app, "model_names", return_value=["m", "other"])
        p.start()
        self.addCleanup(p.stop)

    def stage(self, name="notes.txt", text="staged text"):
        src = self.tmp / name
        src.write_text(text)
        with mock.patch.object(app, "model_capabilities", return_value=[]):
            return app.add_attachments([str(src)], [], "m")[0]

    def test_whitespace_message_without_files_returns_the_same_chat_untouched(self):
        chat = app.empty_chat("m")
        out = app.chat_submit(" \n\t ", chat, "m", [])
        self.assertIs(out[1], chat)
        self.assertEqual(len(out), 9)
        self.assertFalse(app.CHATS_DIR.exists())

    def test_attachment_only_message_is_titled_after_the_file(self):
        pending = self.stage("budget.csv", "a,b\n1,2")
        out = app.chat_submit("", app.empty_chat(), "m", pending)
        chat = out[1]
        self.assertEqual(chat["title"], "budget.csv")
        self.assertEqual(chat["messages"][0]["content"], "")
        stored = Path(chat["messages"][0]["attachments"][0]["path"])
        self.assertEqual(stored.parent, app.ATTACH_DIR / chat["id"])
        self.assertTrue(stored.exists())
        self.assertEqual((out[6], out[7], out[8]["visible"]), ([], "", False))

    def test_markdown_only_message_gets_default_title(self):
        chat = app.chat_submit("**`##`**", app.empty_chat(), "m", [])[1]
        self.assertEqual(chat["title"], "New chat")
        self.assertEqual(chat["messages"][0]["content"], "**`##`**")

    def test_title_of_exactly_48_characters_is_not_shortened(self):
        msg = "a" * 48
        self.assertEqual(app.chat_submit(msg, app.empty_chat(), "m", [])[1]["title"], msg)

    def test_title_of_49_characters_is_cut_with_ellipsis(self):
        msg = "b" * 49
        self.assertEqual(app.chat_submit(msg, app.empty_chat(), "m", [])[1]["title"], "b" * 48 + "...")

    def test_title_cut_drops_the_trailing_space_before_ellipsis(self):
        msg = "x" * 47 + " " + "tail words go here"
        self.assertEqual(app.chat_submit(msg, app.empty_chat(), "m", [])[1]["title"], "x" * 47 + "...")

    def test_message_is_saved_without_surrounding_whitespace(self):
        chat = app.chat_submit("\n   hello there   \n", app.empty_chat(), "m", [])[1]
        self.assertEqual(app.load_chat(chat["id"])["messages"], [{"role": "user", "content": "hello there"}])

    def test_follow_up_in_an_existing_chat_does_not_refresh_the_list(self):
        first = app.chat_submit("one", app.empty_chat(), "m", [])
        second = app.chat_submit("two", first[1], "m", [])
        self.assertEqual(second[1]["id"], first[1]["id"])
        self.assertEqual([m["content"] for m in second[1]["messages"]], ["one", "two"])
        self.assertTrue(is_noop(second[3]) and is_noop(second[4]))
        self.assertIsInstance(first[3], float)
        self.assertEqual(first[4], first[1]["id"])
        self.assertEqual(second[1]["title"], "one")

    def test_submit_does_not_mutate_the_callers_chat(self):
        chat = app.empty_chat("m")
        app.chat_submit("hi", chat, "m", [])
        self.assertEqual((chat["id"], chat["messages"]), (None, []))

    def test_chat_with_messages_but_no_model_takes_the_chosen_model(self):
        chat = make_chat("nm", model=None)
        out = app.chat_submit("continue", chat, "other", [])
        self.assertEqual(out[1]["model"], "other")
        self.assertFalse(out[5]["interactive"])

    def test_no_model_with_pending_files_keeps_the_files_staged(self):
        pending = self.stage()
        out = app.chat_submit("read this", app.empty_chat(), None, pending)
        self.assertIn("Choose a model", out[2][-1]["content"])
        self.assertTrue(Path(pending[0]["path"]).exists())
        self.assertTrue(all(is_noop(u) for u in out[6:]))

    def test_new_chat_id_format_and_unicode_title(self):
        out = app.chat_submit("Привет, как дела?", app.empty_chat(), "m", [])
        chat = out[1]
        self.assertRegex(chat["id"], r"^\d{8}-\d{6}-[0-9a-f]{6}$")
        self.assertEqual(chat["title"], "Привет, как дела?")
        self.assertEqual(out[0], "")
        self.assertEqual(out[2], [{"role": "user", "content": "Привет, как дела?"}])


# ---------------------------------------------------------------- streaming replies
class ChatRespond(TempDataTest):
    def setUp(self):
        super().setUp()
        for name, value in (("scanning_model", None), ("model_capabilities", ["completion"])):
            p = mock.patch.object(app, name, return_value=value)
            p.start()
            self.addCleanup(p.stop)

    def respond(self, chat, lines=(), status=200, side_effect=None):
        resp = FakeResponse(lines=lines, status=status)
        with mock.patch.object(app.requests, "post", return_value=resp, side_effect=side_effect) as post:
            outs = list(app.chat_respond(chat))
        return outs, post

    def last(self, outs):
        return outs[-1][0]["messages"][-1]

    def test_guard_no_chat(self):
        outs, post = self.respond(None)
        self.assertEqual(len(outs), 1)
        self.assertIsNone(outs[0][0])
        post.assert_not_called()

    def test_guard_last_message_already_answered(self):
        chat = make_chat(messages=[{"role": "user", "content": "q"}, {"role": "assistant", "content": "a"}])
        outs, post = self.respond(chat)
        self.assertEqual(len(outs), 1)
        self.assertEqual(len(chat["messages"]), 2)
        post.assert_not_called()

    def test_guard_chat_without_model(self):
        outs, post = self.respond(make_chat(model=None))
        self.assertEqual(len(outs), 1)
        post.assert_not_called()

    def test_http_404_model_not_found_is_reported(self):
        chat = make_chat(model="missing:1b")
        outs, _ = self.respond(chat, [json.dumps({"error": "model 'missing:1b' not found"}).encode()], status=404)
        self.assertEqual(self.last(outs)["content"], "\n\nError: model 'missing:1b' not found")
        self.assertEqual(outs[-1][0]["totals"], app.new_totals())

    # BUG: chat_respond never looks at the HTTP status; a 500 with a plain-text body is parsed as JSON, so the user
    # sees "Error: Expecting value: line 1 column 1 (char 0)" instead of the server failure (app.py chat_respond ~1155).
    def test_http_500_plain_text_body_mentions_the_server_error(self):
        outs, _ = self.respond(make_chat(), [b"Internal Server Error"], status=500)
        self.assertIn("500", self.last(outs)["content"])

    # BUG: an HTTP 503 with an empty body is shown as "*No response.*", as if the model had answered with nothing,
    # instead of an error (app.py chat_respond ~1155: r.status_code / r.ok is never checked).
    def test_http_503_empty_body_is_shown_as_an_error(self):
        outs, _ = self.respond(make_chat(), [], status=503)
        self.assertIn("Error", self.last(outs)["content"])

    def test_error_event_mid_stream_keeps_partial_text_and_stops_reading(self):
        lines = [ev("Partial "), ev("answer"), json.dumps({"error": "out of memory"}).encode(), ev("never shown")]
        outs, _ = self.respond(make_chat(), lines)
        reply = self.last(outs)
        self.assertEqual(reply["content"], "Partial answer\n\nError: out of memory")
        self.assertEqual(outs[-1][0]["totals"]["replies"], 0)

    def test_empty_reply_is_shown_as_no_response_but_still_counted(self):
        outs, _ = self.respond(make_chat(), [ev("", done=True, prompt_eval_count=9, eval_count=0)])
        self.assertEqual(self.last(outs)["content"], "*No response.*")
        self.assertEqual(outs[-1][0]["totals"]["replies"], 1)
        self.assertEqual(outs[-1][0]["totals"]["prompt"], 9)

    def test_thinking_is_hidden_but_counts_as_the_first_token(self):
        lines = [json.dumps({"message": {"role": "assistant", "content": "", "thinking": "Let me think"}}).encode(),
                 ev("42"), ev("", done=True, prompt_eval_count=1, eval_count=1)]
        outs, _ = self.respond(make_chat(), lines)
        reply = self.last(outs)
        self.assertEqual(reply["content"], "42")
        self.assertNotIn("Let me think", reply["content"])
        self.assertIn("to first token", reply["stats"])

    def test_token_usage_and_speed_in_the_stats_line(self):
        lines = [ev("ok"), ev("", done=True, prompt_eval_count=1234, eval_count=20, eval_duration=4e9)]
        outs, _ = self.respond(make_chat(model="qwen3:8b"), lines)
        stats = self.last(outs)["stats"]
        self.assertIn("1,234 prompt + 20 response tokens", stats)
        self.assertIn("5.0 tokens/s", stats)
        self.assertTrue(stats.endswith("qwen3:8b"))
        self.assertEqual(outs[-1][0]["totals"]["prompt"], 1234)

    def test_model_load_time_shown_only_from_half_a_second(self):
        slow, _ = self.respond(make_chat("a"), [ev("x", done=True, load_duration=0.5e9)])
        fast, _ = self.respond(make_chat("b"), [ev("x", done=True, load_duration=0.49e9)])
        self.assertIn("0.5 s model load", self.last(slow)["stats"])
        self.assertNotIn("model load", self.last(fast)["stats"])

    def test_stream_cut_off_before_done_keeps_text_without_counting(self):
        outs, _ = self.respond(make_chat(model="m"), [ev("half an ans"), ev("wer")])
        reply = self.last(outs)
        self.assertEqual(reply["content"], "half an answer")
        self.assertNotIn("tokens", reply["stats"])
        self.assertTrue(reply["stats"].endswith(" &nbsp;·&nbsp; m"))
        self.assertEqual(outs[-1][0]["totals"], app.new_totals())

    def test_garbled_line_mid_stream_becomes_an_error_after_partial_text(self):
        outs, _ = self.respond(make_chat(), [ev("Good start"), b"{garbled", ev("lost")])
        content = self.last(outs)["content"]
        self.assertTrue(content.startswith("Good start\n\nError: "))
        self.assertNotIn("lost", content)

    def test_blank_keepalive_lines_are_ignored(self):
        outs, _ = self.respond(make_chat(), [b"", ev("A"), b"", b"", ev("B", done=True)])
        self.assertEqual(self.last(outs)["content"], "AB")

    def test_request_payload_and_url(self):
        chat = make_chat(model="llama3", messages=[{"role": "user", "content": "q1"},
                                                   {"role": "assistant", "content": "a1", "stats": "s"},
                                                   {"role": "user", "content": "q2"}])
        _, post = self.respond(chat, [ev("a2", done=True)])
        url, kw = post.call_args.args[0], post.call_args.kwargs
        self.assertEqual(url, f"{app.OLLAMA}/api/chat")
        self.assertEqual(kw["json"]["model"], "llama3")
        self.assertEqual(kw["json"]["messages"], [{"role": "user", "content": "q1"},
                                                  {"role": "assistant", "content": "a1"},
                                                  {"role": "user", "content": "q2"}])
        self.assertEqual((kw["json"]["stream"], kw["json"]["keep_alive"], kw["stream"]), (True, "3m", True))

    def test_images_go_to_a_vision_model(self):
        img = self.tmp / "cat.png"
        img.write_bytes(b"PNGDATA")
        chat = make_chat(messages=[{"role": "user", "content": "what is this", "attachments": [
            {"kind": "image", "name": "cat.png", "path": str(img), "size": 7}]}])
        with mock.patch.object(app, "model_capabilities", return_value=["completion", "vision"]) as caps:
            _, post = self.respond(chat, [ev("a cat", done=True)])
        caps.assert_called_once_with("m")
        self.assertEqual(post.call_args.kwargs["json"]["messages"][0]["images"], [base64.b64encode(b"PNGDATA").decode()])
        self.assertNotIn("images not sent", chat["messages"][-1]["stats"])

    def test_images_to_a_text_only_model_are_noted_in_stats(self):
        img = self.tmp / "dog.png"
        img.write_bytes(b"x")
        chat = make_chat(messages=[{"role": "user", "content": "describe", "attachments": [
            {"kind": "image", "name": "dog.png", "path": str(img), "size": 1}]}])
        _, post = self.respond(chat, [ev("cannot see", done=True)])
        sent = post.call_args.kwargs["json"]["messages"][0]
        self.assertNotIn("images", sent)
        self.assertIn("dog.png", sent["content"])
        self.assertIn("images not sent: this model has no vision support", chat["messages"][-1]["stats"])

    def test_totals_accumulate_over_two_replies(self):
        chat = make_chat()
        self.respond(chat, [ev("1", done=True, prompt_eval_count=10, eval_count=3)])
        chat["messages"].append({"role": "user", "content": "more"})
        outs, _ = self.respond(chat, [ev("2", done=True, prompt_eval_count=20, eval_count=4)])
        t = outs[-1][0]["totals"]
        self.assertEqual((t["replies"], t["prompt"], t["output"]), (2, 30, 7))
        self.assertIn("2 replies", outs[-1][4])
        self.assertEqual(app.load_chat("c1")["totals"]["replies"], 2)

    def test_first_update_shows_typing_and_last_restores_buttons(self):
        outs, _ = self.respond(make_chat(), [ev("hey", done=True, eval_count=1)])
        first, final = outs[0], outs[-1]
        self.assertEqual(first[1][-1]["content"], app.TYPING)
        self.assertEqual((first[2]["visible"], first[3]["visible"]), (False, True))
        self.assertEqual((final[2]["visible"], final[3]["visible"]), (True, False))
        self.assertTrue(final[4].startswith("**This chat:** 1 reply"))

    def test_unreadable_model_format_error_gets_advice(self):
        err = 'read GGUF metadata "/home/u/.ollama/models/blobs/sha256-abc": unsupported tensor type PQ2_0'
        outs, _ = self.respond(make_chat(), [json.dumps({"error": err}).encode()], status=500)
        content = self.last(outs)["content"]
        self.assertIn("unsupported tensor type PQ2_0", content)
        self.assertNotIn("/home/u/.ollama", content)
        self.assertIn("format Ollama cannot read", content)

    def test_read_timeout_is_reported(self):
        outs, _ = self.respond(make_chat(), side_effect=app.requests.Timeout("Read timed out."))
        self.assertEqual(self.last(outs)["content"], "\n\nError: Read timed out.")

    def test_partial_text_streams_into_the_chat_and_is_saved_while_streaming(self):
        clock = iter(range(1000, 100000, 3))
        fake_time = types.SimpleNamespace(time=lambda: float(next(clock)), strftime=time.strftime)
        chat = make_chat("live")
        app.save_chat(chat)
        with mock.patch.object(app, "time", fake_time), \
                mock.patch.object(app, "save_chat", wraps=app.save_chat) as save:
            outs, _ = self.respond(chat, [ev("Hel"), ev("lo"), ev("", done=True)])
        shown = [o[1][-1]["content"] for o in outs[1:-1]]
        self.assertEqual(shown, ["Hel", "Hello", "Hello"])
        self.assertTrue(any(c.kwargs.get("index") is False for c in save.call_args_list))
        self.assertEqual(self.last(outs)["content"], "Hello")

    def test_chatting_with_the_scans_model_under_its_short_name_is_allowed(self):
        with mock.patch.object(app, "scanning_model", return_value="llama3:latest"):
            outs, post = self.respond(make_chat(model="llama3"), [ev("fine", done=True)])
        post.assert_called_once()
        self.assertEqual(post.call_args.kwargs["json"]["keep_alive"], -1)
        self.assertEqual(self.last(outs)["content"], "fine")


# ---------------------------------------------------------------- stop, open, new, remove
class ChatLifecycle(TempDataTest):
    def setUp(self):
        super().setUp()
        p = mock.patch.object(app, "model_names", return_value=["m"])
        p.start()
        self.addCleanup(p.stop)

    def test_stop_keeps_partial_reply_and_saves_it(self):
        chat = make_chat("st", messages=[{"role": "user", "content": "q"},
                                         {"role": "assistant", "content": "half done", "stats": ""}])
        send, stop, chat2, shown = app.chat_stopped(chat)
        self.assertEqual((send["visible"], stop["visible"]), (True, False))
        self.assertEqual(chat2["messages"][-1]["content"], "half done")
        self.assertEqual(app.load_chat("st")["messages"][-1]["stats"], "Stopped")
        self.assertEqual(shown[-1]["content"], "half done\n\n<sub>Stopped</sub>")

    def test_stop_before_a_reply_started_changes_nothing(self):
        chat = make_chat("st2")
        out = app.chat_stopped(chat)
        self.assertEqual(out[2]["messages"], [{"role": "user", "content": "hi"}])
        self.assertFalse(app.CHATS_DIR.exists())

    def test_stop_with_no_chat(self):
        out = app.chat_stopped(None)
        self.assertIsNone(out[2])
        self.assertEqual(out[3], [])

    def test_open_missing_chat_explains_and_resets(self):
        chat, shown, lock, totals, _, cid = app.open_chat("20990101-000000-ffffff")
        self.assertEqual((chat["id"], shown, totals, cid), (None, [], "This chat no longer exists.", None))
        self.assertTrue(is_noop(lock))

    def test_open_saved_chat_shows_messages_totals_and_lock(self):
        chat = make_chat("op", messages=[{"role": "user", "content": "q"},
                                         {"role": "assistant", "content": "a", "stats": "1.0 s"}])
        chat["totals"] = {"replies": 1, "seconds": 1.0, "prompt": 5, "output": 6}
        app.save_chat(chat)
        loaded, shown, lock, totals, refresh, cid = app.open_chat("op")
        self.assertEqual(cid, "op")
        self.assertEqual(shown[-1]["content"], "a\n\n<sub>1.0 s</sub>")
        self.assertFalse(lock["interactive"])
        self.assertIn("11 total", totals)
        self.assertIsInstance(refresh, float)

    def test_open_old_chat_without_totals(self):
        chat = make_chat("old")
        del chat["totals"]
        app.save_chat(chat)
        self.assertEqual(app.open_chat("old")[3], "No messages yet.")

    def test_start_new_chat_resets_everything(self):
        chat, shown, totals, refresh, cid, lock = app.start_new_chat()
        self.assertEqual((chat["id"], chat["messages"], shown, totals, cid), (None, [], [], "No messages yet.", None))
        self.assertTrue(lock["interactive"])
        self.assertIsInstance(refresh, float)

    def test_remove_not_confirmed_keeps_everything(self):
        current = make_chat("keep")
        app.save_chat(current)
        out = app.remove_chat("keep", False, current)
        self.assertIs(out[0], current)
        self.assertTrue(all(is_noop(u) for u in out[1:]))
        self.assertIsNotNone(app.load_chat("keep"))

    def test_remove_the_open_chat_starts_a_new_one(self):
        current = make_chat("act")
        app.save_chat(current)
        chat, shown, totals, refresh, cid, lock = app.remove_chat("act", True, current)
        self.assertEqual((chat["id"], shown, totals, cid), (None, [], "No messages yet.", None))
        self.assertTrue(lock["interactive"])
        self.assertIsNone(app.load_chat("act"))

    def test_remove_another_chat_keeps_the_open_one(self):
        current, other = make_chat("cur"), make_chat("oth")
        app.save_chat(current)
        app.save_chat(other)
        (app.ATTACH_DIR / "oth").mkdir(parents=True)
        out = app.remove_chat("oth", True, current)
        self.assertIs(out[0], current)
        self.assertIsInstance(out[3], float)
        self.assertTrue(is_noop(out[1]) and is_noop(out[4]))
        self.assertIsNone(app.load_chat("oth"))
        self.assertFalse((app.ATTACH_DIR / "oth").exists())
        self.assertIsNotNone(app.load_chat("cur"))


# ---------------------------------------------------------------- model capabilities and attach button
class Capabilities(unittest.TestCase):
    def test_capabilities_come_from_api_show(self):
        with mock.patch.object(app.requests, "post",
                               return_value=FakeResponse(data={"capabilities": ["completion", "vision"]})) as post:
            self.assertEqual(app.model_capabilities("gemma3:4b"), ["completion", "vision"])
        self.assertTrue(post.call_args.args[0].endswith("/api/show"))
        self.assertEqual(post.call_args.kwargs["json"], {"model": "gemma3:4b"})

    def test_capabilities_when_ollama_is_down(self):
        with mock.patch.object(app.requests, "post", side_effect=app.requests.ConnectionError("refused")):
            self.assertEqual(app.model_capabilities("m"), [])

    def test_capabilities_missing_or_unparseable(self):
        with mock.patch.object(app.requests, "post", return_value=FakeResponse(data={"error": "model not found"})):
            self.assertEqual(app.model_capabilities("m"), [])
        bad = FakeResponse()
        bad.json = mock.Mock(side_effect=ValueError("not json"))
        with mock.patch.object(app.requests, "post", return_value=bad):
            self.assertEqual(app.model_capabilities("m"), [])

    def test_attach_button_accepts_images_for_vision_models(self):
        with mock.patch.object(app, "model_capabilities", return_value=["vision"]):
            up = app.attach_state("llava")
        self.assertEqual((up["file_types"], up["elem_classes"]), (app.ATTACH_TYPES, ["attach-btn"]))

    def test_attach_button_is_documents_only_for_text_models(self):
        with mock.patch.object(app, "model_capabilities", return_value=["completion", "tools"]):
            up = app.attach_state("qwen3")
        self.assertEqual(up["file_types"], app.DOC_TYPES)
        self.assertIn("docs-only", up["elem_classes"])

    def test_attach_button_without_a_model_asks_nothing(self):
        with mock.patch.object(app, "model_capabilities", side_effect=AssertionError("no model, no lookup")):
            up = app.attach_state(None)
        self.assertEqual(up["file_types"], app.DOC_TYPES)

    def test_attach_types_lists_every_kind_once_and_doc_types_have_no_images(self):
        self.assertEqual(app.ATTACH_TYPES, sorted(set(app.ATTACH_TYPES)))
        for ext in (".png", ".jpg", ".pdf", ".docx", ".txt", ".py", ".csv", ".json"):
            self.assertIn(ext, app.ATTACH_TYPES)
        self.assertEqual(set(app.DOC_TYPES), set(app.ATTACH_TYPES) - app.IMAGE_EXTS)


# ---------------------------------------------------------------- staging attachments
def text_pdf(text):
    """A minimal one-page PDF whose page shows `text` (built with correct xref offsets)."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


class AttachmentScenarios(TempDataTest):
    def upload(self, name, content):
        p = self.tmp / "uploads"
        p.mkdir(exist_ok=True)
        p = p / name
        p.write_bytes(content if isinstance(content, bytes) else content.encode())
        return str(p)

    def add(self, files, pending=None, caps=(), model="m"):
        with mock.patch.object(app, "model_capabilities", return_value=list(caps)):
            return app.add_attachments(files, pending or [], model)

    def staged(self):
        folder = app.ATTACH_DIR / "staging"
        return sorted(p.name for p in folder.iterdir()) if folder.exists() else []

    def test_extract_text_from_docx(self):
        import docx
        d = docx.Document()
        d.add_paragraph("First paragraph")
        d.add_paragraph("Zweiter Absatz ü")
        path = self.tmp / "memo.docx"
        d.save(str(path))
        self.assertEqual(app.extract_text(path), "First paragraph\nZweiter Absatz ü")
        pending, _, _ = self.add([str(path)])
        self.assertEqual(pending[0]["kind"], "text")
        self.assertIn("Zweiter Absatz", pending[0]["text"])

    def test_extract_text_from_pdf(self):
        path = Path(self.upload("paper.pdf", text_pdf("Quantum widgets")))
        self.assertIn("Quantum widgets", app.extract_text(path))
        pending, markup, _ = self.add([str(path)])
        self.assertIn("Quantum widgets", pending[0]["text"])
        self.assertIn("paper.pdf", markup)

    def test_pdf_without_text_is_rejected_and_its_copy_removed(self):
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(width=200, height=200)
        path = self.tmp / "scan.pdf"
        with open(path, "wb") as fh:
            w.write(fh)
        pending, markup, vis = self.add([str(path)])
        self.assertEqual(pending, [])
        self.assertIn("scan.pdf: no readable text (scanned PDFs need OCR)", markup)
        self.assertEqual(self.staged(), [])
        self.assertFalse(vis["visible"])

    def test_corrupt_pdf_cannot_be_read_and_its_copy_is_removed(self):
        pending, markup, _ = self.add([self.upload("broken.pdf", b"this is not a pdf at all")])
        self.assertEqual(pending, [])
        self.assertIn("broken.pdf: could not read (", markup)
        self.assertEqual(self.staged(), [])

    def test_invalid_utf8_is_read_with_replacement_characters(self):
        pending, _, _ = self.add([self.upload("latin1.txt", b"caf\xe9 au lait")])
        self.assertEqual(pending[0]["text"], "caf\ufffd au lait")

    def test_exactly_max_characters_is_not_truncated(self):
        pending, markup, _ = self.add([self.upload("edge.txt", "y" * app.MAX_ATTACH_CHARS)])
        self.assertFalse(pending[0]["truncated"])
        self.assertEqual(len(pending[0]["text"]), app.MAX_ATTACH_CHARS)
        self.assertNotIn("truncated", markup)

    def test_uppercase_extension_is_accepted(self):
        pending, _, _ = self.add([self.upload("README.TXT", "shout")])
        self.assertEqual((pending[0]["name"], pending[0]["text"]), ("README.TXT", "shout"))

    def test_images_need_a_chosen_model(self):
        with mock.patch.object(app, "model_capabilities", side_effect=AssertionError("no lookup without a model")):
            pending, markup, _ = app.add_attachments([self.upload("a.jpg", b"\xff\xd8")], [], None)
        self.assertEqual(pending, [])
        self.assertIn("a.jpg: this model can&#x27;t view images", markup)

    def test_new_files_are_added_after_already_staged_ones(self):
        first, _, _ = self.add([self.upload("one.md", "# one")])
        both, markup, vis = self.add([self.upload("two.log", "line two")], pending=first)
        self.assertEqual([a["name"] for a in both], ["one.md", "two.log"])
        self.assertEqual(len(first), 1)  # the earlier list is not modified in place
        self.assertEqual(markup.count('class="attach-chip"'), 2)
        self.assertTrue(vis["visible"])

    def test_gradio_file_objects_with_a_name_attribute(self):
        upload = types.SimpleNamespace(name=self.upload("obj.txt", "  from an object \n"))
        pending, _, _ = self.add([upload])
        self.assertEqual(pending[0]["text"], "from an object")
        self.assertEqual(pending[0]["size"], len("  from an object \n"))
        self.assertTrue(Path(pending[0]["path"]).is_relative_to(app.ATTACH_DIR / "staging"))

    def test_mixed_batch_keeps_the_good_file_and_reports_the_bad(self):
        files = [self.upload("good.json", '{"a": 1}'), self.upload("bad.bin", b"\x00\x01")]
        pending, markup, vis = self.add(files)
        self.assertEqual([a["name"] for a in pending], ["good.json"])
        self.assertIn("bad.bin: unsupported file type", markup)
        self.assertTrue(vis["visible"])

    def test_images_rejected_when_capabilities_cannot_be_checked(self):
        with mock.patch.object(app.requests, "post", side_effect=app.requests.ConnectionError("down")):
            pending, markup, _ = app.add_attachments([self.upload("p.webp", b"RIFF")], [], "llava")
        self.assertEqual(pending, [])
        self.assertIn("can&#x27;t view images", markup)

    def test_chips_show_image_thumbnail_and_megabytes(self):
        markup = app.attachments_html([{"name": "big.png", "kind": "image", "size": 3_400_000, "path": "/s/big.png"}])
        self.assertIn('<img src="/gradio_api/file=/s/big.png" alt="">', markup)
        self.assertIn("3.4 MB", markup)
        self.assertNotIn("chip-doc", markup)

    def test_errors_without_files_show_only_error_lines(self):
        markup = app.attachments_html([], ["x.exe: unsupported file type", "y.png: <bad>"])
        self.assertNotIn("attach-chips", markup)
        self.assertEqual(markup.count('class="note-line error"'), 2)
        self.assertIn("&lt;bad&gt;", markup)
        self.assertEqual(app.attachments_html(None), "")

    def test_clear_attachments_deletes_staged_files(self):
        pending, _, _ = self.add([self.upload("a.txt", "a"), self.upload("b.txt", "b")])
        self.assertEqual(len(self.staged()), 2)
        out = app.clear_attachments(pending)
        self.assertEqual((out[0], out[1], out[2]["visible"]), ([], "", False))
        self.assertEqual(self.staged(), [])

    def test_clear_attachments_tolerates_missing_files_and_none(self):
        app.clear_attachments([{"path": str(self.tmp / "already-gone.txt")}])
        self.assertEqual(app.clear_attachments(None)[0], [])

    def test_attach_to_chat_skips_vanished_files_and_keeps_details(self):
        pending, _, _ = self.add([self.upload("keep.txt", "keep"), self.upload("lost.txt", "lost")])
        Path(pending[1]["path"]).unlink()
        stored = app._attach_to_chat(pending, "chat-9")
        self.assertEqual([a["name"] for a in stored], ["keep.txt"])
        self.assertEqual(Path(stored[0]["path"]).parent, app.ATTACH_DIR / "chat-9")
        self.assertEqual((stored[0]["text"], stored[0]["kind"]), ("keep", "text"))

    def test_attachment_markup_for_an_image_escapes_its_name(self):
        markup = app.attachment_markup([{"kind": "image", "name": 'a"><script>.png', "path": "/c/a.png", "size": 1}])
        self.assertIn('class="attach-img" src="/gradio_api/file=/c/a.png"', markup)
        self.assertNotIn("<script>", markup)
        self.assertEqual(app.attachment_markup(None), "")
        self.assertEqual(app.attachment_markup([]), "")

    def test_attachment_markup_for_a_document_is_an_in_message_chip(self):
        markup = app.attachment_markup([{"kind": "text", "name": "spec.md", "path": "/c/spec.md", "size": 1_500_000}])
        self.assertIn('class="attach-chip in-message"', markup)
        self.assertIn("spec.md", markup)
        self.assertIn("1.5 MB", markup)

    def test_api_messages_mark_truncated_files(self):
        chat = {"messages": [{"role": "user", "content": "summarise", "attachments": [
            {"kind": "text", "name": "huge.log", "text": "TOP", "truncated": True}]}]}
        msgs, _ = app.api_messages(chat, vision=False)
        self.assertEqual(msgs[0]["content"], "summarise\n\n[Attached file: huge.log (truncated)]\nTOP")

    def test_api_messages_describe_images_to_text_only_models(self):
        chat = {"messages": [{"role": "user", "content": "hi", "attachments": [
            {"kind": "image", "name": "a.png", "path": "/nope/a.png"}, {"kind": "image", "name": "b.gif", "path": "/nope/b.gif"}]}]}
        msgs, skipped = app.api_messages(chat, vision=False)
        self.assertEqual(skipped, ["a.png", "b.gif"])
        self.assertIn("[An image named a.png was attached, but this model cannot view images.]", msgs[0]["content"])
        self.assertIn("b.gif", msgs[0]["content"])

    def test_api_messages_flatten_list_content_and_keep_roles(self):
        chat = {"messages": [{"role": "user", "content": [{"text": "part one "}, "part two"]},
                             {"role": "assistant", "content": "reply", "stats": "1 s", "model": "m"}]}
        msgs, skipped = app.api_messages(chat, vision=True)
        self.assertEqual(msgs, [{"role": "user", "content": "part one part two"}, {"role": "assistant", "content": "reply"}])
        self.assertEqual(skipped, [])

    def test_api_messages_send_several_images_in_order(self):
        paths = []
        for name, data in (("1.png", b"one"), ("2.png", b"two")):
            p = self.tmp / name
            p.write_bytes(data)
            paths.append({"kind": "image", "name": name, "path": str(p)})
        msgs, _ = app.api_messages({"messages": [{"role": "user", "content": "compare", "attachments": paths}]}, vision=True)
        self.assertEqual(msgs[0]["images"], [base64.b64encode(b"one").decode(), base64.b64encode(b"two").decode()])
        self.assertEqual(msgs[0]["content"], "compare")


TEXT_EXT_CASES = [
    (".csv", "name,score\nalice,10"), (".json", '{"key": "value"}'), (".py", "def f():\n    return 1"),
    (".yaml", "a: 1\nb: [2, 3]"), (".sql", "SELECT * FROM t;"), (".html", "<p>hello</p>"),
]
UNSUPPORTED_CASES = [("archive.zip", b"PK\x03\x04"), ("song.mp3", b"ID3"), ("logo.svg", b"<svg/>"), ("Makefile", b"all:")]


for _ext, _body in TEXT_EXT_CASES:
    def _t(self, ext=_ext, body=_body):
        pending, markup, _ = self.add([self.upload(f"file{ext}", body)])
        self.assertEqual((pending[0]["kind"], pending[0]["text"], pending[0]["truncated"]), ("text", body, False))
        msgs, _ = app.api_messages({"messages": [{"role": "user", "content": "q", "attachments": pending}]}, vision=False)
        self.assertEqual(msgs[0]["content"], f"q\n\n[Attached file: file{ext}]\n{body}")
    setattr(AttachmentScenarios, f"test_text_attachment_type_{_ext[1:]}_is_inlined", _t)

for _fname, _data in UNSUPPORTED_CASES:
    def _t(self, fname=_fname, data=_data):
        pending, markup, vis = self.add([self.upload(fname, data)], caps=["vision"])
        self.assertEqual(pending, [])
        self.assertIn(f"{fname}: unsupported file type", markup)
        self.assertEqual(self.staged(), [])
    setattr(AttachmentScenarios, f"test_unsupported_upload_{re.sub(r'[^a-z0-9]', '_', _fname.lower())}_is_rejected", _t)


if __name__ == "__main__":
    unittest.main()
